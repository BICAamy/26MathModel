#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
2026 高教社杯 A题：药材的烘干问题 —— 问题1
第二套独立数值算法：有限差分法(FDM) + 显式 FTCS 时间推进

与原程序的区别：
    原程序：有限体积法(FVM) + solve_ivp(RK45)
    本程序：节点有限差分(FDM) + Forward Euler / FTCS + Robin 边界幽灵点

模型保持不变：
    1) 一维圆柱径向非稳态热传导
    2) 非线性 Fick 水分扩散
    3) 中心零梯度（对称边界）
    4) 表面对流换热 / 对流传质（Robin 边界）
    5) 附件1的烘房温度、水分浓度做分段线性插值

默认输入：
    file/附件1.xlsx

默认输出：
    output/q1/result1_fdm.xlsx

推荐交叉验证：
    python q1/q1_solver_fdm.py \
        --input file/附件1.xlsx \
        --output output/q1/result1_fdm.xlsx \
        --reference output/q1/result1.xlsx

更细网格验证（可选）：
    python q1/q1_solver_fdm.py \
        --input file/附件1.xlsx \
        --output output/q1/result1_fdm_refined.xlsx \
        --reference output/q1/result1.xlsx \
        --dr-cm 0.025 --dt 0.05
"""

from pathlib import Path
import argparse
import numpy as np
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment


# ============================================================
# 1. 读取附件1
# ============================================================

def read_attachment1(path: Path):
    """读取附件1前3列：时间/s、烘房温度/°C、烘房水分浓度/(kg/kg)。"""
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise ValueError("附件1没有有效数据。")

    data = []
    for row in rows[1:]:
        if row[0] is None:
            continue
        try:
            t = float(row[0])
            T = float(row[1])
            C = float(row[2])
            data.append((t, T, C))
        except (TypeError, ValueError):
            continue

    if not data:
        raise ValueError("附件1中没有读取到有效数值数据。")

    arr = np.asarray(data, dtype=float)
    return arr[:, 0], arr[:, 1], arr[:, 2]


# ============================================================
# 2. 第二套算法：FDM + FTCS
# ============================================================

def solve_q1_fdm(input_path: Path, dr_cm=0.1, dt=0.25):
    # ---------------- 题目给定参数 ----------------
    rho = 820.0             # kg/m^3
    cp = 2600.0             # J/(kg·K)
    k = 0.36                # W/(m·K)
    h = 25.0                # W/(m^2·K)
    hm = 8.0e-7             # m/s

    R = 0.02                # m = 2 cm
    dr = dr_cm / 100.0      # cm -> m
    alpha = k / (rho * cp)

    N_float = R / dr
    N = int(round(N_float))
    if abs(N - N_float) > 1e-10:
        raise ValueError("dr 必须能整除 2 cm 半径，例如 0.1、0.05、0.025 cm。")

    r = np.arange(N + 1, dtype=float) * dr

    # FTCS 对中心节点的热传导稳定性条件近似：4*alpha*dt/dr^2 <= 1
    dt_heat_limit = dr**2 / (4.0 * alpha)
    if dt > dt_heat_limit:
        raise ValueError(
            f"当前 dt={dt:g}s 对 dr={dr_cm:g}cm 过大，显式格式可能不稳定。"
            f"建议 dt <= {dt_heat_limit:.6g}s。"
        )

    # 为了严格每 1 s 保存一次，要求 1/dt 为整数
    steps_per_second = int(round(1.0 / dt))
    if not np.isclose(steps_per_second * dt, 1.0, atol=1e-12):
        raise ValueError("为了每1秒准确输出，dt 应能整除 1 s，例如 0.5、0.25、0.2、0.1、0.05 s。")

    # ---------------- 读取附件1 ----------------
    time_data, Tair_data, Cair_data = read_attachment1(input_path)
    mask = (time_data >= 0.0) & (time_data <= 1800.0)
    time_q1 = time_data[mask]
    Tair_q1 = Tair_data[mask]
    Cair_q1 = Cair_data[mask]

    if len(time_q1) < 2 or time_q1[0] > 0.0 or time_q1[-1] < 1800.0:
        raise ValueError("附件1必须完整覆盖 0~1800 s。")

    def T_air(t):
        return float(np.interp(t, time_q1, Tair_q1))

    def C_air(t):
        return float(np.interp(t, time_q1, Cair_q1))

    # D(C) = 7e-9 * exp(-0.89/C)
    def D_of_C(C):
        C_safe = np.maximum(np.asarray(C, dtype=float), 1e-12)
        return 7.0e-9 * np.exp(-0.89 / C_safe)

    # D'(C) = D(C) * 0.89/C^2
    def Dprime_of_C(C):
        C_safe = np.maximum(np.asarray(C, dtype=float), 1e-12)
        D = D_of_C(C_safe)
        return D * 0.89 / (C_safe**2)

    # ---------------- 初值 ----------------
    T = np.full(N + 1, 28.0, dtype=float)
    C = np.full(N + 1, 2.55, dtype=float)

    # 每1秒保存一次，先保存 t=0
    T_hist = np.empty((1801, N + 1), dtype=float)
    C_hist = np.empty((1801, N + 1), dtype=float)
    T_hist[0] = T
    C_hist[0] = C

    total_steps = int(round(1800.0 / dt))
    out_index = 1

    print("开始用 FDM + FTCS 求解……")
    print(f"空间步长 dr = {dr_cm:g} cm，共 {N+1} 个节点")
    print(f"时间步长 dt = {dt:g} s，共 {total_steps} 个时间步")

    for n in range(total_steps):
        t = n * dt
        Tinf = T_air(t)
        Cinf = C_air(t)

        dTdt = np.zeros_like(T)
        dCdt = np.zeros_like(C)

        # ====================================================
        # A. 温度场：
        # dT/dt = alpha * (T_rr + (1/r) T_r)
        # ====================================================

        # ---- 中心 r=0：利用极限，圆柱径向 Laplacian -> 2*T_rr ----
        dTdt[0] = 4.0 * alpha * (T[1] - T[0]) / dr**2

        # ---- 内部节点 1...N-1：二阶中心差分 ----
        if N > 1:
            ri = r[1:N]
            T_r = (T[2:] - T[:-2]) / (2.0 * dr)
            T_rr = (T[2:] - 2.0 * T[1:N] + T[:-2]) / dr**2
            dTdt[1:N] = alpha * (T_rr + T_r / ri)

        # ---- 表面 r=R：Robin 边界 + 幽灵点 ----
        # -k*T_r(R) = h*(T_R - T_air)
        # => T_r(R) = -(h/k)*(T_R-T_air)
        # 用中心差分构造外部幽灵点 T_{N+1}
        T_ghost = T[N - 1] - 2.0 * dr * (h / k) * (T[N] - Tinf)
        T_r_R = (T_ghost - T[N - 1]) / (2.0 * dr)
        T_rr_R = (T_ghost - 2.0 * T[N] + T[N - 1]) / dr**2
        dTdt[N] = alpha * (T_rr_R + T_r_R / R)

        # ====================================================
        # B. 水分场：
        # dC/dt = (1/r) d/dr [r D(C) C_r]
        #       = D(C)*(C_rr + C_r/r) + D'(C)*(C_r)^2
        # ====================================================

        # ---- 中心 r=0：C_r=0，因此 D'(C)*(C_r)^2 项为0 ----
        dCdt[0] = 4.0 * D_of_C(C[0]) * (C[1] - C[0]) / dr**2

        # ---- 内部节点：二阶中心差分，直接离散展开后的 PDE ----
        if N > 1:
            ri = r[1:N]
            C_r = (C[2:] - C[:-2]) / (2.0 * dr)
            C_rr = (C[2:] - 2.0 * C[1:N] + C[:-2]) / dr**2
            D_i = D_of_C(C[1:N])
            Dp_i = Dprime_of_C(C[1:N])

            dCdt[1:N] = (
                D_i * (C_rr + C_r / ri)
                + Dp_i * (C_r**2)
            )

        # ---- 表面 r=R：Robin 边界 + 幽灵点 ----
        # -D(C_R)*C_r(R) = hm*(C_R-C_air)
        D_R = float(D_of_C(C[N]))
        Dp_R = float(Dprime_of_C(C[N]))

        C_r_bc = -(hm / D_R) * (C[N] - Cinf)
        C_ghost = C[N - 1] + 2.0 * dr * C_r_bc

        C_r_R = (C_ghost - C[N - 1]) / (2.0 * dr)
        C_rr_R = (C_ghost - 2.0 * C[N] + C[N - 1]) / dr**2

        dCdt[N] = (
            D_R * (C_rr_R + C_r_R / R)
            + Dp_R * (C_r_R**2)
        )

        # ====================================================
        # C. Forward Euler / FTCS 时间推进
        # ====================================================
        T = T + dt * dTdt
        C = C + dt * dCdt

        # 每 1 s 保存一次
        if (n + 1) % steps_per_second == 0:
            T_hist[out_index] = T
            C_hist[out_index] = C
            out_index += 1

    print("FDM + FTCS 求解完成。")

    times = np.arange(0, 1801, dtype=int)
    r_cm = r * 100.0
    return times, r_cm, T_hist, C_hist


# ============================================================
# 3. 将求解网格映射到题目要求的 0.1 cm 输出网格
# ============================================================

def resample_to_required_grid(r_cm, field):
    target_r_cm = np.round(np.arange(0.0, 2.0 + 0.1, 0.1), 10)
    out = np.empty((field.shape[0], target_r_cm.size), dtype=float)

    for i in range(field.shape[0]):
        out[i] = np.interp(target_r_cm, r_cm, field[i])

    return target_r_cm, out


# ============================================================
# 4. 写 result1_fdm.xlsx
# ============================================================

def write_result1(output_path: Path, times, r_cm, T, C):
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws_T = wb.active
    ws_T.title = "温度"
    ws_C = wb.create_sheet("水分浓度")

    headers = ["时间/s"] + [round(float(x), 1) for x in r_cm]
    ws_T.append(headers)
    ws_C.append(headers)

    for i, t in enumerate(times):
        ws_T.append([int(t)] + [float(v) for v in T[i]])
        ws_C.append([int(t)] + [float(v) for v in C[i]])

    fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    white_font = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center")

    for ws in (ws_T, ws_C):
        for cell in ws[1]:
            cell.fill = fill
            cell.font = white_font
            cell.alignment = center

        ws.freeze_panes = "B2"
        ws.column_dimensions["A"].width = 11

        for col in range(2, 23):
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 9

        for row in ws.iter_rows(min_row=2, min_col=2):
            for cell in row:
                cell.number_format = "0.0000"
                cell.alignment = center

        for cell in ws["A"]:
            cell.alignment = center

    wb.save(output_path)


# ============================================================
# 5. 打印题目要求表1、表2
# ============================================================

def print_required_tables(T, C):
    report_times = [100, 300, 600, 900, 1200, 1500, 1800]
    report_r_cm = [0.0, 0.5, 1.0, 1.5, 2.0]
    idx = [int(round(x / 0.1)) for x in report_r_cm]

    print("\n==============================")
    print("FDM 表1：30分钟内药材的温度 / °C")
    print("==============================")
    print("时间/s | " + " | ".join(f"{x:g}cm" for x in report_r_cm))

    for t in report_times:
        print(
            f"{t:6d} | "
            + " | ".join(f"{T[t, j]:.4f}" for j in idx)
        )

    print("\n======================================")
    print("FDM 表2：30分钟内药材的水分浓度 / kg/kg")
    print("======================================")
    print("时间/s | " + " | ".join(f"{x:g}cm" for x in report_r_cm))

    for t in report_times:
        print(
            f"{t:6d} | "
            + " | ".join(f"{C[t, j]:.4f}" for j in idx)
        )


# ============================================================
# 6. 读取参考 result1.xlsx 并比较
# ============================================================

def read_result1(path: Path):
    wb = load_workbook(path, data_only=True)

    if "温度" not in wb.sheetnames or "水分浓度" not in wb.sheetnames:
        raise ValueError("参考文件必须包含“温度”和“水分浓度”两个工作表。")

    def read_sheet(ws):
        rows = list(ws.iter_rows(values_only=True))
        r_cm = np.asarray([float(v) for v in rows[0][1:] if v is not None])
        times = []
        data = []

        for row in rows[1:]:
            if row[0] is None:
                continue
            vals = row[1:1 + len(r_cm)]
            if any(v is None for v in vals):
                continue
            times.append(int(round(float(row[0]))))
            data.append([float(v) for v in vals])

        return np.asarray(times, dtype=int), r_cm, np.asarray(data, dtype=float)

    tT, rT, T = read_sheet(wb["温度"])
    tC, rC, C = read_sheet(wb["水分浓度"])

    if not np.array_equal(tT, tC) or not np.allclose(rT, rC):
        raise ValueError("参考文件两个工作表的时间/空间网格不一致。")

    return tT, rT, T, C


def compare_with_reference(reference_path: Path, times, r_cm, T, C):
    ref_t, ref_r, T_ref, C_ref = read_result1(reference_path)

    # 当前程序输出固定为题目要求的 0.1 cm 网格，因此应能直接对齐
    if not np.array_equal(times, ref_t):
        raise ValueError("参考结果的时间网格与当前结果不一致。")
    if not np.allclose(r_cm, ref_r, atol=1e-10):
        raise ValueError("参考结果的空间网格与当前结果不一致。")

    dT = T - T_ref
    dC = C - C_ref

    print("\n======================================")
    print("与原 FVM + RK45 结果的全场误差")
    print("======================================")
    print(f"温度最大绝对误差 MAE_max = {np.max(np.abs(dT)):.6g} °C")
    print(f"温度平均绝对误差 MAE     = {np.mean(np.abs(dT)):.6g} °C")
    print(f"水分最大绝对误差 MAE_max = {np.max(np.abs(dC)):.6g} kg/kg")
    print(f"水分平均绝对误差 MAE     = {np.mean(np.abs(dC)):.6g} kg/kg")

    # 再比较论文要求的 7×5 关键点
    report_times = [100, 300, 600, 900, 1200, 1500, 1800]
    report_r_cm = [0.0, 0.5, 1.0, 1.5, 2.0]
    idx_r = [int(round(x / 0.1)) for x in report_r_cm]

    max_T_key = 0.0
    max_C_key = 0.0
    for t in report_times:
        max_T_key = max(max_T_key, np.max(np.abs(dT[t, idx_r])))
        max_C_key = max(max_C_key, np.max(np.abs(dC[t, idx_r])))

    print("\n论文表1/表2关键点：")
    print(f"温度关键点最大绝对差 = {max_T_key:.6g} °C")
    print(f"水分关键点最大绝对差 = {max_C_key:.6g} kg/kg")


# ============================================================
# 7. 主程序
# ============================================================

def main():
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--input",
        default="file/附件1.xlsx",
        help="附件1.xlsx路径"
    )
    parser.add_argument(
        "--output",
        default="output/q1/result1_fdm.xlsx",
        help="第二算法结果输出路径"
    )
    parser.add_argument(
        "--reference",
        default=None,
        help="可选：原 FVM+RK45 的 result1.xlsx，用于自动误差比较"
    )
    parser.add_argument(
        "--dr-cm",
        type=float,
        default=0.1,
        help="FDM 求解网格空间步长/cm，默认0.1"
    )
    parser.add_argument(
        "--dt",
        type=float,
        default=0.25,
        help="FTCS时间步长/s，默认0.25"
    )

    args = parser.parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)

    if not input_path.exists():
        raise FileNotFoundError(f"找不到输入文件：{input_path.resolve()}")

    times, r_solver_cm, T_solver, C_solver = solve_q1_fdm(
        input_path,
        dr_cm=args.dr_cm,
        dt=args.dt,
    )

    # 无论内部求解网格多细，最终统一输出题目要求的 0.1 cm 网格
    r_out_cm, T_out = resample_to_required_grid(r_solver_cm, T_solver)
    _, C_out = resample_to_required_grid(r_solver_cm, C_solver)

    print_required_tables(T_out, C_out)
    write_result1(output_path, times, r_out_cm, T_out, C_out)

    print(f"\n第二算法结果文件：{output_path.resolve()}")

    if args.reference is not None:
        reference_path = Path(args.reference)
        if not reference_path.exists():
            raise FileNotFoundError(f"找不到参考文件：{reference_path.resolve()}")
        compare_with_reference(
            reference_path,
            times,
            r_out_cm,
            T_out,
            C_out,
        )


if __name__ == "__main__":
    main()
