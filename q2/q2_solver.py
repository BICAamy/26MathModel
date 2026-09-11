#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
2026 高教社杯 A 题“药材的烘干问题”——问题 2 数值求解器

数值方法
--------
1. 在固定半径圆柱的径向上采用守恒型有限体积法（FVM）；
2. 变物性系数 k(C)、D(C,T) 在相邻控制体界面取调和平均；
3. 温度场和水分场联立，形成非线性刚性常微分方程组；
4. 使用 scipy.integrate.solve_ivp 的 BDF 隐式方法推进时间；
5. 烘房温度与水分浓度由附件 1 分段线性插值获得；
6. 输出 0~3 h 每隔 1 s、半径方向每隔 0.1 cm 的 result2.xlsx。

默认从项目根目录运行：

python q2/q2_solver.py \
    --input file/附件1.xlsx \
    --output output/q2/result2.xlsx

依赖：numpy、scipy、openpyxl

说明
----
附录 3 只重新给出了 rho、cp、k、D 的经验公式，没有另给表面对流
换热/传质系数。因此本程序按建模稿沿用附录 2 的
h = 25 W/(m^2 K)、hm = 8e-7 m/s；两者均可通过命令行参数修改。
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path
from time import perf_counter

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from scipy.integrate import solve_ivp
from scipy.sparse import lil_matrix


@dataclass(frozen=True)
class ModelConfig:
    """问题 2 的几何、边界与数值参数。"""

    radius_m: float = 0.02
    n_intervals: int = 200
    t_end_s: float = 3.0 * 3600.0
    output_dt_s: float = 1.0
    initial_temperature_c: float = 28.0
    initial_moisture: float = 2.55
    heat_transfer_coefficient: float = 25.0
    mass_transfer_coefficient: float = 8.0e-7
    method: str = "BDF"
    rtol: float = 1.0e-7
    atol_temperature: float = 1.0e-8
    atol_moisture: float = 1.0e-10
    max_step_s: float = 30.0

    def validate(self) -> None:
        if self.radius_m <= 0:
            raise ValueError("radius_m 必须为正数。")
        if self.n_intervals < 4:
            raise ValueError("n_intervals 至少为 4。")
        if self.t_end_s <= 0 or self.output_dt_s <= 0:
            raise ValueError("t_end_s 和 output_dt_s 必须为正数。")
        if self.heat_transfer_coefficient <= 0:
            raise ValueError("对流换热系数必须为正数。")
        if self.mass_transfer_coefficient <= 0:
            raise ValueError("对流传质系数必须为正数。")
        if self.initial_moisture <= 0:
            raise ValueError("初始水分浓度必须为正数。")


@dataclass
class EnvironmentData:
    """附件 1 中的烘房边界数据。"""

    time_s: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray

    def values(self, t: float) -> tuple[float, float]:
        """对附件 1 做分段线性插值，返回 T_inf(t)、C_inf(t)。"""

        temperature = float(np.interp(t, self.time_s, self.temperature_c))
        moisture = float(np.interp(t, self.time_s, self.moisture))
        return temperature, moisture


@dataclass
class SimulationResult:
    """数值解与用于守恒检查的网格信息。"""

    time_s: np.ndarray
    radius_nodes_m: np.ndarray
    temperature_nodes_c: np.ndarray
    moisture_nodes: np.ndarray
    control_volumes_m3_per_m: np.ndarray
    nfev: int
    njev: int
    nlu: int
    elapsed_s: float


def _as_float(value: object, cell_name: str) -> float:
    """把 Excel 单元格转换为有限浮点数，并给出可定位的报错。"""

    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{cell_name} 不是有效数值：{value!r}") from exc
    if not np.isfinite(number):
        raise ValueError(f"{cell_name} 不是有限数值：{value!r}")
    return number


def read_environment(path: Path, required_end_s: float) -> EnvironmentData:
    """读取附件 1 的前三列：时间/s、温度/°C、水分浓度/(kg/kg)。"""

    if not path.exists():
        raise FileNotFoundError(f"找不到附件 1：{path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]

    rows: list[tuple[float, float, float]] = []
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=2, min_col=1, max_col=3, values_only=True),
        start=2,
    ):
        if all(value is None for value in row):
            continue
        if any(value is None for value in row):
            raise ValueError(f"附件 1 第 {row_index} 行前三列存在空值：{row!r}")
        rows.append(
            (
                _as_float(row[0], f"A{row_index}"),
                _as_float(row[1], f"B{row_index}"),
                _as_float(row[2], f"C{row_index}"),
            )
        )
    workbook.close()

    if len(rows) < 2:
        raise ValueError("附件 1 至少需要两个有效数据点。")

    data = np.asarray(rows, dtype=float)
    order = np.argsort(data[:, 0])
    data = data[order]
    time_s, temperature_c, moisture = data.T

    if np.any(np.diff(time_s) <= 0):
        raise ValueError("附件 1 的时间列必须严格递增，且不能有重复时间。")
    if time_s[0] > 0:
        raise ValueError("附件 1 必须覆盖 t=0。")
    if time_s[-1] + 1.0e-9 < required_end_s:
        raise ValueError(
            f"附件 1 只覆盖到 {time_s[-1]:g} s，"
            f"不足以计算到 {required_end_s:g} s；程序不会无依据外推。"
        )
    if np.any(temperature_c <= -273.15):
        raise ValueError("附件 1 中存在不合法温度（低于绝对零度）。")
    if np.any(moisture < 0):
        raise ValueError("附件 1 中存在负的水分浓度。")

    return EnvironmentData(time_s, temperature_c, moisture)


def density(moisture: np.ndarray) -> np.ndarray:
    """rho(C) = 650 + 128 C，单位 kg/m^3。"""

    return 650.0 + 128.0 * moisture


def specific_heat(moisture: np.ndarray) -> np.ndarray:
    """cp(C) = 1450 + 2736 C/(C+1)，单位 J/(kg K)。"""

    return 1450.0 + 2736.0 * moisture / (moisture + 1.0)


def thermal_conductivity(moisture: np.ndarray) -> np.ndarray:
    """k(C) = 0.21 + 0.38 C/(C+1)，单位 W/(m K)。"""

    return 0.21 + 0.38 * moisture / (moisture + 1.0)


def moisture_diffusivity(
    moisture: np.ndarray,
    temperature_c: np.ndarray,
) -> np.ndarray:
    """附录 3 的 D(C,T)，其中经验式中的 T 必须使用 K。"""

    # 正常物理解在前三小时内远离 0；下限仅防止隐式迭代试探点出现除零。
    moisture_safe = np.maximum(moisture, 1.0e-10)
    temperature_k = np.maximum(temperature_c + 273.15, 1.0)
    exponent = -0.45 / moisture_safe - 3850.0 / temperature_k
    return 2.4e-3 * np.exp(np.clip(exponent, -745.0, 50.0))


def harmonic_mean(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    """正系数的调和平均，用于控制体公共界面。"""

    denominator = left + right
    return np.divide(
        2.0 * left * right,
        denominator,
        out=np.zeros_like(denominator),
        where=denominator > 0,
    )


def build_radial_grid(
    radius_m: float,
    n_intervals: int,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, float]:
    """
    构造包含 r=0 和 r=R 的节点型有限体积网格。

    每个节点对应一个控制体；首、末节点对应半控制体。几何量按单位
    圆柱长度计算，长度因子在通量与体积中相消。
    """

    radius_nodes = np.linspace(0.0, radius_m, n_intervals + 1)
    dr = radius_m / n_intervals
    n_nodes = radius_nodes.size

    radius_faces = np.empty(n_nodes + 1, dtype=float)
    radius_faces[0] = 0.0
    radius_faces[-1] = radius_m
    radius_faces[1:-1] = 0.5 * (radius_nodes[:-1] + radius_nodes[1:])

    face_areas = 2.0 * np.pi * radius_faces
    control_volumes = np.pi * (
        radius_faces[1:] ** 2 - radius_faces[:-1] ** 2
    )
    return radius_nodes, face_areas, control_volumes, dr


def build_jacobian_sparsity(n_nodes: int):
    """给 BDF/Radau 提供温度—水分局部耦合的稀疏结构。"""

    pattern = lil_matrix((2 * n_nodes, 2 * n_nodes), dtype=np.int8)
    for i in range(n_nodes):
        for j in range(max(0, i - 1), min(n_nodes, i + 2)):
            pattern[i, j] = 1
            pattern[i, n_nodes + j] = 1
            pattern[n_nodes + i, j] = 1
            pattern[n_nodes + i, n_nodes + j] = 1
    return pattern.tocsr()


def make_rhs(
    config: ModelConfig,
    environment: EnvironmentData,
    face_areas: np.ndarray,
    control_volumes: np.ndarray,
    dr: float,
):
    """构造半离散系统 dy/dt=f(t,y)。"""

    n_nodes = control_volumes.size
    heat_face_rate = np.empty(n_nodes + 1, dtype=float)
    moisture_face_rate = np.empty(n_nodes + 1, dtype=float)

    def rhs(t: float, state: np.ndarray) -> np.ndarray:
        temperature = state[:n_nodes]
        moisture = state[n_nodes:]

        # 隐式求解器可能在牛顿迭代中短暂试探到极小负值。仅在物性公式中
        # 做数值保护，不改写状态变量本身。
        moisture_for_properties = np.maximum(moisture, 1.0e-10)

        rho = density(moisture_for_properties)
        cp = specific_heat(moisture_for_properties)
        k_node = thermal_conductivity(moisture_for_properties)
        d_node = moisture_diffusivity(moisture_for_properties, temperature)

        k_face = harmonic_mean(k_node[:-1], k_node[1:])
        d_face = harmonic_mean(d_node[:-1], d_node[1:])

        # 统一把“沿半径向外”的传递率记为正。
        # 中心对称边界：r=0 处面积为 0，净通量为 0。
        heat_face_rate[0] = 0.0
        moisture_face_rate[0] = 0.0

        heat_face_rate[1:n_nodes] = (
            -face_areas[1:n_nodes]
            * k_face
            * np.diff(temperature)
            / dr
        )
        moisture_face_rate[1:n_nodes] = (
            -face_areas[1:n_nodes]
            * d_face
            * np.diff(moisture)
            / dr
        )

        ambient_temperature, ambient_moisture = environment.values(t)
        heat_face_rate[-1] = (
            face_areas[-1]
            * config.heat_transfer_coefficient
            * (temperature[-1] - ambient_temperature)
        )
        moisture_face_rate[-1] = (
            face_areas[-1]
            * config.mass_transfer_coefficient
            * (moisture[-1] - ambient_moisture)
        )

        # 对第 i 个控制体：积累量 = 西侧向外率 - 东侧向外率。
        net_heat_rate = heat_face_rate[:-1] - heat_face_rate[1:]
        net_moisture_rate = moisture_face_rate[:-1] - moisture_face_rate[1:]

        d_temperature_dt = net_heat_rate / (rho * cp * control_volumes)
        d_moisture_dt = net_moisture_rate / control_volumes
        return np.concatenate((d_temperature_dt, d_moisture_dt))

    return rhs


def make_output_times(t_end_s: float, output_dt_s: float) -> np.ndarray:
    """生成包含 t=0 与 t=t_end 的输出时刻。"""

    n_regular = int(np.floor(t_end_s / output_dt_s + 1.0e-12))
    times = output_dt_s * np.arange(n_regular + 1, dtype=float)
    if times[-1] < t_end_s - 1.0e-9:
        times = np.append(times, t_end_s)
    else:
        times[-1] = t_end_s
    return times


def solve_problem2(
    environment: EnvironmentData,
    config: ModelConfig,
) -> SimulationResult:
    """联立求解问题 2 的温度场和水分场。"""

    config.validate()
    radius_nodes, face_areas, control_volumes, dr = build_radial_grid(
        config.radius_m,
        config.n_intervals,
    )
    n_nodes = radius_nodes.size

    initial_state = np.concatenate(
        (
            np.full(n_nodes, config.initial_temperature_c, dtype=float),
            np.full(n_nodes, config.initial_moisture, dtype=float),
        )
    )
    output_times = make_output_times(config.t_end_s, config.output_dt_s)
    absolute_tolerance = np.concatenate(
        (
            np.full(n_nodes, config.atol_temperature, dtype=float),
            np.full(n_nodes, config.atol_moisture, dtype=float),
        )
    )

    rhs = make_rhs(config, environment, face_areas, control_volumes, dr)
    jacobian_sparsity = build_jacobian_sparsity(n_nodes)

    started = perf_counter()
    solution = solve_ivp(
        rhs,
        (0.0, config.t_end_s),
        initial_state,
        method=config.method,
        t_eval=output_times,
        rtol=config.rtol,
        atol=absolute_tolerance,
        max_step=config.max_step_s,
        jac_sparsity=jacobian_sparsity,
    )
    elapsed = perf_counter() - started

    if not solution.success:
        raise RuntimeError(f"时间积分失败：{solution.message}")
    if not np.all(np.isfinite(solution.y)):
        raise RuntimeError("数值解中出现 NaN 或 Inf。")

    temperature_nodes = solution.y[:n_nodes, :].T
    moisture_nodes = solution.y[n_nodes:, :].T
    if float(np.min(moisture_nodes)) <= 0:
        raise RuntimeError(
            "数值解出现非正水分浓度；请加密网格、减小 max_step 或检查参数。"
        )

    return SimulationResult(
        time_s=solution.t,
        radius_nodes_m=radius_nodes,
        temperature_nodes_c=temperature_nodes,
        moisture_nodes=moisture_nodes,
        control_volumes_m3_per_m=control_volumes,
        nfev=solution.nfev,
        njev=solution.njev,
        nlu=solution.nlu,
        elapsed_s=elapsed,
    )


def interpolate_to_output_radii(
    field: np.ndarray,
    radius_nodes_m: np.ndarray,
    output_radius_cm: np.ndarray,
) -> np.ndarray:
    """把所有时刻的节点解线性插值到规定输出半径。"""

    output_radius_m = output_radius_cm / 100.0
    if output_radius_m[0] < -1.0e-12:
        raise ValueError("输出半径不能为负。")
    if output_radius_m[-1] > radius_nodes_m[-1] + 1.0e-12:
        raise ValueError("输出半径超出药材半径。")

    right = np.searchsorted(radius_nodes_m, output_radius_m, side="right")
    left = np.clip(right - 1, 0, radius_nodes_m.size - 2)
    right = left + 1
    span = radius_nodes_m[right] - radius_nodes_m[left]
    weight_right = (output_radius_m - radius_nodes_m[left]) / span

    return (
        field[:, left] * (1.0 - weight_right)[None, :]
        + field[:, right] * weight_right[None, :]
    )


def required_output_radii_cm(radius_m: float) -> np.ndarray:
    """题目要求半径方向每隔 0.1 cm 输出。"""

    radius_cm = 100.0 * radius_m
    count = int(round(radius_cm / 0.1))
    if not np.isclose(count * 0.1, radius_cm, atol=1.0e-10):
        raise ValueError("当前半径不能被 0.1 cm 整除，无法生成题目规定网格。")
    return np.linspace(0.0, radius_cm, count + 1)


def _styled_header_cell(worksheet, value: object) -> WriteOnlyCell:
    cell = WriteOnlyCell(worksheet, value=value)
    cell.font = Font(bold=True)
    cell.number_format = "0.0"
    return cell


def _four_decimal_cell(worksheet, value: float) -> WriteOnlyCell:
    cell = WriteOnlyCell(worksheet, value=round(float(value), 4))
    cell.number_format = "0.0000"
    return cell


def write_result2(
    output_path: Path,
    time_s: np.ndarray,
    output_radius_cm: np.ndarray,
    temperature_c: np.ndarray,
    moisture: np.ndarray,
) -> None:
    """按附件 3 说明生成 result2.xlsx。"""

    if temperature_c.shape != moisture.shape:
        raise ValueError("温度矩阵与水分矩阵形状不一致。")
    expected_shape = (time_s.size, output_radius_cm.size)
    if temperature_c.shape != expected_shape:
        raise ValueError(
            f"结果矩阵形状应为 {expected_shape}，实际为 {temperature_c.shape}。"
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    temperature_sheet = workbook.create_sheet("温度")
    moisture_sheet = workbook.create_sheet("水分浓度")

    for worksheet in (temperature_sheet, moisture_sheet):
        worksheet.freeze_panes = "B2"
        worksheet.column_dimensions["A"].width = 12
        header = [_styled_header_cell(worksheet, "时间/s")]
        header.extend(
            _styled_header_cell(worksheet, round(float(radius), 1))
            for radius in output_radius_cm
        )
        worksheet.append(header)

    for index, time_value in enumerate(time_s):
        if np.isclose(time_value, round(time_value), atol=1.0e-9):
            time_cell: float | int = int(round(time_value))
        else:
            time_cell = float(time_value)

        temperature_sheet.append(
            [time_cell]
            + [
                _four_decimal_cell(temperature_sheet, value)
                for value in temperature_c[index]
            ]
        )
        moisture_sheet.append(
            [time_cell]
            + [
                _four_decimal_cell(moisture_sheet, value)
                for value in moisture[index]
            ]
        )

    temporary_path = output_path.with_name(output_path.stem + ".tmp.xlsx")
    workbook.save(temporary_path)
    temporary_path.replace(output_path)


def value_at_time(
    time_s: np.ndarray,
    field: np.ndarray,
    query_time_s: float,
) -> np.ndarray:
    """对全部半径位置做一次时间线性插值。"""

    if query_time_s < time_s[0] or query_time_s > time_s[-1]:
        raise ValueError("查询时刻超出求解区间。")
    right = int(np.searchsorted(time_s, query_time_s, side="left"))
    if right == 0:
        return field[0].copy()
    if right == time_s.size:
        return field[-1].copy()
    if np.isclose(time_s[right], query_time_s, atol=1.0e-10):
        return field[right].copy()
    left = right - 1
    weight = (query_time_s - time_s[left]) / (time_s[right] - time_s[left])
    return (1.0 - weight) * field[left] + weight * field[right]


def print_requested_tables(
    time_s: np.ndarray,
    output_radius_cm: np.ndarray,
    temperature_c: np.ndarray,
    moisture: np.ndarray,
) -> None:
    """打印论文表 3、表 4 所需的 6×5 数值。"""

    latest_requested_time_h = min(3.0, float(time_s[-1]) / 3600.0)
    requested_times_h = np.arange(0.5, latest_requested_time_h + 0.25, 0.5)
    requested_radius_cm = np.asarray([0.0, 0.5, 1.0, 1.5, 2.0])
    radius_indices = [
        int(np.argmin(np.abs(output_radius_cm - radius)))
        for radius in requested_radius_cm
    ]

    def print_one(title: str, field: np.ndarray) -> None:
        print(f"\n{title}")
        print("时间/h | " + " | ".join(f"{r:g} cm" for r in requested_radius_cm))
        if requested_times_h.size == 0:
            print("当前求解时长不足 0.5 h，暂无题目规定的汇总时刻。")
            return
        for time_h in requested_times_h:
            values = value_at_time(time_s, field, time_h * 3600.0)
            formatted = " | ".join(f"{values[i]:.4f}" for i in radius_indices)
            print(f"{time_h:>5.1f} | {formatted}")

    print_one("表 3：3 小时内药材的温度/°C", temperature_c)
    print_one("表 4：3 小时内药材的水分浓度/(kg/kg)", moisture)


def print_diagnostics(
    result: SimulationResult,
    environment: EnvironmentData,
) -> None:
    """输出基本物理检查，帮助识别符号、单位或数值错误。"""

    volume_weights = result.control_volumes_m3_per_m
    average_moisture = (
        result.moisture_nodes @ volume_weights / np.sum(volume_weights)
    )
    largest_average_increase = float(np.max(np.diff(average_moisture)))

    ambient_temperature_on_output = np.interp(
        result.time_s,
        environment.time_s,
        environment.temperature_c,
    )
    maximum_principle_excess = float(
        np.max(result.temperature_nodes_c)
        - max(
            float(np.max(ambient_temperature_on_output)),
            float(result.temperature_nodes_c[0, 0]),
        )
    )

    print("\n数值诊断")
    print(
        f"求解耗时 {result.elapsed_s:.3f} s；"
        f"nfev={result.nfev}, njev={result.njev}, nlu={result.nlu}"
    )
    print(
        "温度范围："
        f"[{np.min(result.temperature_nodes_c):.6f}, "
        f"{np.max(result.temperature_nodes_c):.6f}] °C"
    )
    print(
        "水分浓度范围："
        f"[{np.min(result.moisture_nodes):.6f}, "
        f"{np.max(result.moisture_nodes):.6f}] kg/kg"
    )
    print(
        "体积加权平均水分："
        f"{average_moisture[0]:.6f} -> {average_moisture[-1]:.6f} kg/kg"
    )
    print(f"平均水分单步最大回升量：{largest_average_increase:.3e}")
    print(f"温度最大值原理超限量：{maximum_principle_excess:.3e} °C")

    tolerance = 5.0e-8
    if largest_average_increase > tolerance:
        raise RuntimeError("平均水分出现异常回升，请检查边界符号或数值精度。")
    if maximum_principle_excess > 1.0e-5:
        raise RuntimeError("温度超过初值/边界上界，请检查边界符号或数值精度。")


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="求解 2026 高教社杯 A 题问题 2，并生成 result2.xlsx。"
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=Path("file/附件1.xlsx"),
        help="附件 1 路径（默认：file/附件1.xlsx）",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("output/result2.xlsx"),
        help="结果文件路径（默认：output/result2.xlsx）",
    )
    parser.add_argument("--radius-m", type=float, default=0.02)
    parser.add_argument("--n-intervals", type=int, default=200)
    parser.add_argument("--t-end-s", type=float, default=10800.0)
    parser.add_argument("--output-dt-s", type=float, default=1.0)
    parser.add_argument("--initial-temperature-c", type=float, default=28.0)
    parser.add_argument("--initial-moisture", type=float, default=2.55)
    parser.add_argument("--h-heat", type=float, default=25.0)
    parser.add_argument("--h-mass", type=float, default=8.0e-7)
    parser.add_argument("--method", choices=("BDF", "Radau"), default="BDF")
    parser.add_argument("--rtol", type=float, default=1.0e-7)
    parser.add_argument("--atol-temperature", type=float, default=1.0e-8)
    parser.add_argument("--atol-moisture", type=float, default=1.0e-10)
    parser.add_argument("--max-step-s", type=float, default=30.0)
    parser.add_argument(
        "--skip-excel",
        action="store_true",
        help="只求解并打印结果，不写 Excel（用于网格收敛测试）。",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_arguments()
    config = ModelConfig(
        radius_m=args.radius_m,
        n_intervals=args.n_intervals,
        t_end_s=args.t_end_s,
        output_dt_s=args.output_dt_s,
        initial_temperature_c=args.initial_temperature_c,
        initial_moisture=args.initial_moisture,
        heat_transfer_coefficient=args.h_heat,
        mass_transfer_coefficient=args.h_mass,
        method=args.method,
        rtol=args.rtol,
        atol_temperature=args.atol_temperature,
        atol_moisture=args.atol_moisture,
        max_step_s=args.max_step_s,
    )

    environment = read_environment(args.input, config.t_end_s)
    result = solve_problem2(environment, config)

    output_radius_cm = required_output_radii_cm(config.radius_m)
    temperature_output = interpolate_to_output_radii(
        result.temperature_nodes_c,
        result.radius_nodes_m,
        output_radius_cm,
    )
    moisture_output = interpolate_to_output_radii(
        result.moisture_nodes,
        result.radius_nodes_m,
        output_radius_cm,
    )

    print_requested_tables(
        result.time_s,
        output_radius_cm,
        temperature_output,
        moisture_output,
    )
    print_diagnostics(result, environment)

    if not args.skip_excel:
        write_result2(
            args.output,
            result.time_s,
            output_radius_cm,
            temperature_output,
            moisture_output,
        )
        print(f"\n已生成：{args.output.resolve()}")


if __name__ == "__main__":
    main()
