#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
2026 高教社杯 A 题“药材的烘干问题”——问题 2 论文绘图脚本

一次运行生成 3 组核心论文图，每张图包含 2 个子图：

图1：温度场 / 水分浓度场二维时空分布图
图2：中心—表面温差 / 水分浓度差随时间变化图
图3：热导率 k / 有效水分扩散系数 D 的时空演化图

推荐放置位置：
    26MathModel/
    ├── output/q2/result2.xlsx
    └── plot/
        ├── q1/q1_plot.py
        └── q2/q2_plot.py   <- 本脚本

从项目根目录运行：
    python plot/q2/q2_plot.py

默认输出：
    plot/output/q2/
        q2_fig1_spatiotemporal_fields.png
        q2_fig1_spatiotemporal_fields.pdf
        q2_fig2_center_surface_differences.png
        q2_fig2_center_surface_differences.pdf
        q2_fig3_variable_properties.png
        q2_fig3_variable_properties.pdf

说明：
1. 三张图均直接读取 output/q2/result2.xlsx，不重复求解 PDE。
2. k(C)、D(C,T) 与 q2/q2_solver.py 中附录 3 的经验公式保持一致。
3. D(C,T) 中温度必须使用 K，因此由 result2 的摄氏温度统一加 273.15。
4. 所有图同时保存 600 dpi PNG 与 PDF，适合直接插入 Word 论文。
5. 子图编号 (a)、(b) 统一放在子图下方。
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Tuple

import numpy as np
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.ticker import MaxNLocator
from openpyxl import load_workbook


# ============================================================
# 0. 路径与默认配置
# ============================================================

SCRIPT_DIR = Path(__file__).resolve().parent

# 文件固定放在 26MathModel/plot/q2/q2_plot.py：
# SCRIPT_DIR = .../26MathModel/plot/q2
# parents[0] = .../plot
# parents[1] = .../26MathModel
PROJECT_ROOT = SCRIPT_DIR.parents[1]

DEFAULT_RESULT = "output/q2/result2.xlsx"
DEFAULT_OUTPUT_DIR = "plot/output/q2"


# ============================================================
# 1. 路径与绘图风格
# ============================================================

def resolve_input_path(path_str: str) -> Path:
    """依次尝试当前工作目录、项目根目录、脚本目录。"""
    p = Path(path_str).expanduser()
    if p.is_absolute():
        return p

    candidates = [
        Path.cwd() / p,
        PROJECT_ROOT / p,
        SCRIPT_DIR / p,
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate.resolve()

    return (PROJECT_ROOT / p).resolve()


def resolve_output_path(path_str: str) -> Path:
    """相对输出路径统一以项目根目录为基准。"""
    p = Path(path_str).expanduser()
    if p.is_absolute():
        return p
    return (PROJECT_ROOT / p).resolve()


def configure_matplotlib() -> None:
    """与问题 1 绘图脚本保持接近的中文论文风格。"""
    available_fonts = {f.name for f in font_manager.fontManager.ttflist}
    candidates = [
        "PingFang SC",          # macOS
        "Microsoft YaHei",     # Windows
        "SimHei",              # Windows / Linux 常见中文字体
        "Noto Sans CJK SC",    # Linux
        "Source Han Sans SC",  # 思源黑体
        "Arial Unicode MS",
        "DejaVu Sans",
    ]

    for font_name in candidates:
        if font_name in available_fonts:
            plt.rcParams["font.sans-serif"] = [font_name, "DejaVu Sans"]
            break

    plt.rcParams.update(
        {
            "axes.unicode_minus": False,
            "font.size": 10.5,
            "axes.titlesize": 11.5,
            "axes.labelsize": 11,
            "xtick.labelsize": 9.5,
            "ytick.labelsize": 9.5,
            "legend.fontsize": 9.5,
            "figure.dpi": 120,
            "savefig.dpi": 600,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "mathtext.fontset": "stix",
        }
    )


def style_axis(ax: plt.Axes) -> None:
    """统一折线图坐标轴风格。"""
    ax.grid(True, linestyle="--", linewidth=0.55, alpha=0.35)
    ax.tick_params(direction="in", top=True, right=True)
    for spine in ax.spines.values():
        spine.set_linewidth(0.8)


def add_panel_labels(axes: np.ndarray | list[plt.Axes]) -> None:
    """把 (a)、(b)、... 放在各子图下方。"""
    axes_flat = np.asarray(axes, dtype=object).ravel()
    for i, ax in enumerate(axes_flat):
        label = f"({chr(ord('a') + i)})"
        ax.text(
            0.5,
            -0.245,
            label,
            transform=ax.transAxes,
            ha="center",
            va="top",
            fontsize=11.0,
            clip_on=False,
        )


def save_figure(fig: plt.Figure, output_dir: Path, stem: str) -> None:
    """同时保存 600 dpi PNG 与 PDF。"""
    output_dir.mkdir(parents=True, exist_ok=True)
    png_path = output_dir / f"{stem}.png"
    pdf_path = output_dir / f"{stem}.pdf"

    fig.savefig(png_path, dpi=600, bbox_inches="tight", facecolor="white")
    fig.savefig(pdf_path, bbox_inches="tight", facecolor="white")

    print(f"  PNG: {png_path}")
    print(f"  PDF: {pdf_path}")


# ============================================================
# 2. 读取 result2.xlsx
# ============================================================

def _read_field_sheet(ws) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    """
    读取题目模板形式工作表：
        A 列：时间/s
        第 1 行 B...：到中心距离/cm
        数据区域：场变量
    """
    rows = ws.iter_rows(values_only=True)

    try:
        header = next(rows)
    except StopIteration as exc:
        raise ValueError(f"工作表 {ws.title!r} 为空。") from exc

    r_values = []
    for value in header[1:]:
        if value is None:
            break
        r_values.append(float(value))

    r_cm = np.asarray(r_values, dtype=float)
    if r_cm.size == 0:
        raise ValueError(f"工作表 {ws.title!r} 第 1 行没有读取到空间网格。")

    times: list[float] = []
    data: list[list[float]] = []
    n_r = r_cm.size

    for row in rows:
        if row[0] is None:
            continue

        values = row[1: 1 + n_r]
        if len(values) < n_r or any(v is None for v in values):
            raise ValueError(
                f"工作表 {ws.title!r} 在时间 {row[0]!r} 附近存在缺失数据。"
            )

        times.append(float(row[0]))
        data.append([float(v) for v in values])

    if not data:
        raise ValueError(f"工作表 {ws.title!r} 没有读取到有效数据矩阵。")

    return (
        np.asarray(times, dtype=float),
        r_cm,
        np.asarray(data, dtype=float),
    )


def read_result2(path: Path) -> Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """读取 result2.xlsx 中温度场与水分浓度场。"""
    if not path.exists():
        raise FileNotFoundError(f"找不到结果文件：{path}")

    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        required = {"温度", "水分浓度"}
        if not required.issubset(set(wb.sheetnames)):
            raise ValueError(
                f"{path.name} 必须包含“温度”和“水分浓度”两个工作表；"
                f"当前工作表为：{wb.sheetnames}"
            )

        t_T, r_T, T = _read_field_sheet(wb["温度"])
        t_C, r_C, C = _read_field_sheet(wb["水分浓度"])
    finally:
        wb.close()

    if t_T.shape != t_C.shape or not np.allclose(t_T, t_C, atol=1e-12):
        raise ValueError("温度表和水分浓度表的时间网格不一致。")
    if r_T.shape != r_C.shape or not np.allclose(r_T, r_C, atol=1e-12):
        raise ValueError("温度表和水分浓度表的空间网格不一致。")
    if T.shape != C.shape:
        raise ValueError("温度场与水分浓度场矩阵维度不一致。")
    if T.shape != (t_T.size, r_T.size):
        raise ValueError("结果矩阵维度与时间/空间网格不匹配。")
    if not np.all(np.isfinite(T)) or not np.all(np.isfinite(C)):
        raise ValueError("结果中存在 NaN 或 Inf。")
    if np.any(C <= 0):
        raise ValueError("结果中存在非正水分浓度，无法计算 D(C,T)。")

    # 问题 2 正常输出应覆盖 0~3 h、r=0~2 cm。
    if t_T[0] > 1e-9 or t_T[-1] < 10800 - 1e-9:
        raise ValueError("result2.xlsx 必须至少覆盖 0~10800 s。")
    if r_T[0] > 1e-12 or r_T[-1] < 2.0 - 1e-12:
        raise ValueError("result2.xlsx 必须覆盖 r=0~2 cm。")

    return t_T, r_T, T, C


# ============================================================
# 3. 与 q2_solver.py 完全一致的变物性经验公式
# ============================================================

def thermal_conductivity(moisture: np.ndarray) -> np.ndarray:
    """k(C) = 0.21 + 0.38 C/(C+1)，单位 W/(m·K)。"""
    return 0.21 + 0.38 * moisture / (moisture + 1.0)


def moisture_diffusivity(
    moisture: np.ndarray,
    temperature_c: np.ndarray,
) -> np.ndarray:
    """D(C,T)，温度必须转为 K，单位 m²/s。"""
    moisture_safe = np.maximum(moisture, 1.0e-10)
    temperature_k = np.maximum(temperature_c + 273.15, 1.0)
    exponent = -0.45 / moisture_safe - 3850.0 / temperature_k
    return 2.4e-3 * np.exp(np.clip(exponent, -745.0, 50.0))


# ============================================================
# 4. 图1：温度场 / 水分浓度场二维时空分布
# ============================================================

def plot_spatiotemporal_fields(
    times_s: np.ndarray,
    r_cm: np.ndarray,
    T: np.ndarray,
    C: np.ndarray,
    output_dir: Path,
) -> None:
    t_h = times_s / 3600.0

    fig, axes = plt.subplots(1, 2, figsize=(12.2, 4.85))
    fig.subplots_adjust(left=0.075, right=0.96, bottom=0.24, top=0.89, wspace=0.34)

    # (a) 温度场
    levels_T = np.linspace(float(np.min(T)), float(np.max(T)), 64)
    cf_T = axes[0].contourf(
        t_h,
        r_cm,
        T.T,
        levels=levels_T,
        cmap="magma",
        extend="both",
    )
    contour_T = axes[0].contour(
        t_h,
        r_cm,
        T.T,
        levels=np.linspace(float(np.min(T)), float(np.max(T)), 9),
        colors="black",
        linewidths=0.32,
        alpha=0.32,
    )
    axes[0].clabel(contour_T, inline=True, fontsize=7.2, fmt="%.1f")
    cbar_T = fig.colorbar(cf_T, ax=axes[0], pad=0.025)
    cbar_T.set_label(r"温度 $T/{}^\circ\mathrm{C}$")

    axes[0].set_title("药材内部温度场")
    axes[0].set_xlabel(r"时间 $t/\mathrm{h}$")
    axes[0].set_ylabel(r"到药材中心距离 $r/\mathrm{cm}$")
    axes[0].set_xlim(0.0, 3.0)
    axes[0].set_ylim(r_cm[0], r_cm[-1])
    axes[0].set_xticks(np.arange(0.0, 3.01, 0.5))
    axes[0].set_yticks(np.arange(0.0, 2.01, 0.5))
    axes[0].tick_params(direction="in")

    # (b) 水分场
    levels_C = np.linspace(float(np.min(C)), float(np.max(C)), 64)
    cf_C = axes[1].contourf(
        t_h,
        r_cm,
        C.T,
        levels=levels_C,
        cmap="viridis",
        extend="both",
    )
    contour_C = axes[1].contour(
        t_h,
        r_cm,
        C.T,
        levels=np.linspace(float(np.min(C)), float(np.max(C)), 9),
        colors="black",
        linewidths=0.32,
        alpha=0.32,
    )
    axes[1].clabel(contour_C, inline=True, fontsize=7.2, fmt="%.2f")
    cbar_C = fig.colorbar(cf_C, ax=axes[1], pad=0.025)
    cbar_C.set_label(r"水分浓度 $C/(\mathrm{kg}\cdot\mathrm{kg}^{-1})$")

    axes[1].set_title("药材内部水分浓度场")
    axes[1].set_xlabel(r"时间 $t/\mathrm{h}$")
    axes[1].set_ylabel(r"到药材中心距离 $r/\mathrm{cm}$")
    axes[1].set_xlim(0.0, 3.0)
    axes[1].set_ylim(r_cm[0], r_cm[-1])
    axes[1].set_xticks(np.arange(0.0, 3.01, 0.5))
    axes[1].set_yticks(np.arange(0.0, 2.01, 0.5))
    axes[1].tick_params(direction="in")

    add_panel_labels(axes)

    print("保存图1：温度场 / 水分浓度场二维时空分布")
    save_figure(fig, output_dir, "q2_fig1_spatiotemporal_fields")
    plt.close(fig)


# ============================================================
# 5. 图2：中心—表面差值随时间变化
# ============================================================

def _annotate_peak(
    ax: plt.Axes,
    x: np.ndarray,
    y: np.ndarray,
    unit: str,
    value_digits: int,
) -> tuple[float, float]:
    i = int(np.argmax(y))
    x_peak = float(x[i])
    y_peak = float(y[i])

    line_color = ax.lines[-1].get_color()
    ax.scatter(x_peak, y_peak, s=28, color=line_color, zorder=5)
    ax.annotate(
        f"最大值 {y_peak:.{value_digits}f} {unit}\n$t={x_peak:.3f}$ h",
        xy=(x_peak, y_peak),
        xytext=(14, -36),
        textcoords="offset points",
        fontsize=8.8,
        arrowprops=dict(arrowstyle="->", linewidth=0.7),
        bbox=dict(
            boxstyle="round,pad=0.25",
            facecolor="white",
            edgecolor="0.75",
            alpha=0.88,
        ),
    )
    return x_peak, y_peak


def plot_center_surface_differences(
    times_s: np.ndarray,
    r_cm: np.ndarray,
    T: np.ndarray,
    C: np.ndarray,
    output_dir: Path,
) -> None:
    t_h = times_s / 3600.0

    i_center = int(np.argmin(np.abs(r_cm - 0.0)))
    i_surface = int(np.argmax(r_cm))

    # 温度：表面通常高于中心，因此定义为 T_surface - T_center。
    delta_T = T[:, i_surface] - T[:, i_center]

    # 水分：中心通常高于表面，因此定义为 C_center - C_surface。
    delta_C = C[:, i_center] - C[:, i_surface]

    fig, axes = plt.subplots(1, 2, figsize=(12.2, 4.85))
    fig.subplots_adjust(left=0.08, right=0.97, bottom=0.24, top=0.89, wspace=0.30)

    axes[0].plot(t_h, delta_T, linewidth=1.65)
    peak_t_T, peak_delta_T = _annotate_peak(
        axes[0], t_h, delta_T, "°C", value_digits=4
    )
    axes[0].set_title("中心—表面径向温差")
    axes[0].set_xlabel(r"时间 $t/\mathrm{h}$")
    axes[0].set_ylabel(r"径向温差 $\Delta T/{}^\circ\mathrm{C}$")
    axes[0].set_xlim(0.0, 3.0)
    axes[0].set_xticks(np.arange(0.0, 3.01, 0.5))
    axes[0].set_ylim(bottom=0.0)
    axes[0].yaxis.set_major_locator(MaxNLocator(6))
    style_axis(axes[0])

    axes[1].plot(t_h, delta_C, linewidth=1.65)
    peak_t_C, peak_delta_C = _annotate_peak(
        axes[1], t_h, delta_C, "kg/kg", value_digits=4
    )
    axes[1].set_title("中心—表面水分浓度差")
    axes[1].set_xlabel(r"时间 $t/\mathrm{h}$")
    axes[1].set_ylabel(
        r"径向水分浓度差 $\Delta C/(\mathrm{kg}\cdot\mathrm{kg}^{-1})$"
    )
    axes[1].set_xlim(0.0, 3.0)
    axes[1].set_xticks(np.arange(0.0, 3.01, 0.5))
    axes[1].set_ylim(bottom=0.0)
    axes[1].yaxis.set_major_locator(MaxNLocator(6))
    style_axis(axes[1])

    add_panel_labels(axes)

    print("保存图2：中心—表面差值随时间变化")
    save_figure(fig, output_dir, "q2_fig2_center_surface_differences")
    plt.close(fig)

    print("\n========== 图2特征量 ==========")
    print(
        f"最大径向温差：{peak_delta_T:.6f} °C，"
        f"发生于 t={peak_t_T:.6f} h"
    )
    print(
        f"3 h 径向温差：{delta_T[-1]:.6f} °C"
    )
    print(
        f"最大径向水分浓度差：{peak_delta_C:.6f} kg/kg，"
        f"发生于 t={peak_t_C:.6f} h"
    )
    print(
        f"3 h 径向水分浓度差：{delta_C[-1]:.6f} kg/kg"
    )


# ============================================================
# 6. 图3：变物性 + 热质耦合参数时空演化
# ============================================================

def plot_variable_properties(
    times_s: np.ndarray,
    r_cm: np.ndarray,
    T: np.ndarray,
    C: np.ndarray,
    output_dir: Path,
) -> None:
    t_h = times_s / 3600.0

    # 与 q2_solver.py 保持一致。
    k = thermal_conductivity(C)
    D = moisture_diffusivity(C, T)

    # D 数值约为 10^-9~10^-8 m²/s，乘 1e9 后色条更易读。
    D_nano = D * 1.0e9

    fig, axes = plt.subplots(1, 2, figsize=(12.2, 4.85))
    fig.subplots_adjust(left=0.075, right=0.96, bottom=0.24, top=0.89, wspace=0.34)

    # (a) k(C)
    levels_k = np.linspace(float(np.min(k)), float(np.max(k)), 64)
    cf_k = axes[0].contourf(
        t_h,
        r_cm,
        k.T,
        levels=levels_k,
        cmap="cividis",
        extend="both",
    )
    contour_k = axes[0].contour(
        t_h,
        r_cm,
        k.T,
        levels=np.linspace(float(np.min(k)), float(np.max(k)), 8),
        colors="black",
        linewidths=0.32,
        alpha=0.30,
    )
    axes[0].clabel(contour_k, inline=True, fontsize=7.2, fmt="%.3f")
    cbar_k = fig.colorbar(cf_k, ax=axes[0], pad=0.025)
    cbar_k.set_label(r"热导率 $k/(\mathrm{W}\cdot\mathrm{m}^{-1}\cdot\mathrm{K}^{-1})$")

    axes[0].set_title(r"变物性热导率 $k(C)$")
    axes[0].set_xlabel(r"时间 $t/\mathrm{h}$")
    axes[0].set_ylabel(r"到药材中心距离 $r/\mathrm{cm}$")
    axes[0].set_xlim(0.0, 3.0)
    axes[0].set_ylim(r_cm[0], r_cm[-1])
    axes[0].set_xticks(np.arange(0.0, 3.01, 0.5))
    axes[0].set_yticks(np.arange(0.0, 2.01, 0.5))
    axes[0].tick_params(direction="in")

    # (b) D(C,T)：同时受温度和含水率影响，直接体现热—质耦合。
    levels_D = np.linspace(float(np.min(D_nano)), float(np.max(D_nano)), 64)
    cf_D = axes[1].contourf(
        t_h,
        r_cm,
        D_nano.T,
        levels=levels_D,
        cmap="plasma",
        extend="both",
    )
    contour_D = axes[1].contour(
        t_h,
        r_cm,
        D_nano.T,
        levels=np.linspace(float(np.min(D_nano)), float(np.max(D_nano)), 8),
        colors="black",
        linewidths=0.32,
        alpha=0.30,
    )
    axes[1].clabel(contour_D, inline=True, fontsize=7.2, fmt="%.2f")
    cbar_D = fig.colorbar(cf_D, ax=axes[1], pad=0.025)
    cbar_D.set_label(r"有效扩散系数 $D/(10^{-9}\,\mathrm{m^2\,s^{-1}})$")

    axes[1].set_title(r"热—质耦合扩散系数 $D(C,T)$")
    axes[1].set_xlabel(r"时间 $t/\mathrm{h}$")
    axes[1].set_ylabel(r"到药材中心距离 $r/\mathrm{cm}$")
    axes[1].set_xlim(0.0, 3.0)
    axes[1].set_ylim(r_cm[0], r_cm[-1])
    axes[1].set_xticks(np.arange(0.0, 3.01, 0.5))
    axes[1].set_yticks(np.arange(0.0, 2.01, 0.5))
    axes[1].tick_params(direction="in")

    add_panel_labels(axes)

    print("保存图3：变物性参数 k(C) / D(C,T) 时空演化")
    save_figure(fig, output_dir, "q2_fig3_variable_properties")
    plt.close(fig)

    print("\n========== 图3参数范围 ==========")
    print(
        f"k 范围：{np.min(k):.8f} ~ {np.max(k):.8f} W/(m·K)"
    )
    print(
        f"D 范围：{np.min(D):.8e} ~ {np.max(D):.8e} m²/s"
    )


# ============================================================
# 7. 主程序
# ============================================================

def main() -> None:
    parser = argparse.ArgumentParser(
        description="问题2论文绘图：生成时空场、中心—表面差值、变物性耦合参数 3 张图。"
    )
    parser.add_argument(
        "--result",
        default=DEFAULT_RESULT,
        help=f"问题2结果文件，默认：{DEFAULT_RESULT}",
    )
    parser.add_argument(
        "--output-dir",
        default=DEFAULT_OUTPUT_DIR,
        help=f"图片输出目录，默认：{DEFAULT_OUTPUT_DIR}",
    )
    args = parser.parse_args()

    configure_matplotlib()

    result_path = resolve_input_path(args.result)
    output_dir = resolve_output_path(args.output_dir)

    print("========== 问题2论文绘图 ==========")
    print(f"项目根目录：{PROJECT_ROOT}")
    print(f"结果文件  ：{result_path}")
    print(f"图片目录  ：{output_dir}")

    times_s, r_cm, T, C = read_result2(result_path)

    print("\n读取 result2.xlsx 完成")
    print(f"  T/C shape：{T.shape}")
    print(f"  时间范围 ：{times_s[0]:g} ~ {times_s[-1]:g} s")
    print(f"  空间范围 ：{r_cm[0]:g} ~ {r_cm[-1]:g} cm")
    print(f"  温度范围 ：{np.min(T):.4f} ~ {np.max(T):.4f} °C")
    print(f"  水分范围 ：{np.min(C):.4f} ~ {np.max(C):.4f} kg/kg")

    plot_spatiotemporal_fields(times_s, r_cm, T, C, output_dir)
    plot_center_surface_differences(times_s, r_cm, T, C, output_dir)
    plot_variable_properties(times_s, r_cm, T, C, output_dir)

    print("\n全部完成：共生成 3 张论文图，每张含 2 个子图。")


if __name__ == "__main__":
    main()
