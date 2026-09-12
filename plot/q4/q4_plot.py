#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
2026 高教社杯 A 题“药材的烘干问题”——问题 4 论文绘图脚本

一次运行生成 3 张问题 4 核心论文图：

图1：药材半径 R(t) 随时间变化（附件2实测点 + PCHIP 插值）
图2：移动边界下水分浓度时空云图（叠加 R(t) 与 C=0.15 等值线）
图3：中心/内部/表面水分浓度随时间变化（叠加 C=0.15 与烘干结束时刻）

推荐放置位置：
    26MathModel/
    ├── file/附件2.xlsx
    ├── output/q4/result4.xlsx
    └── plot/
        └── q4/
            └── q4_plot.py   <- 本脚本

从项目根目录运行：
    python plot/q4/q4_plot.py

默认输出：
    plot/output/q4/
        q4_fig1_radius_shrinkage.png/.pdf
        q4_fig2_moving_boundary_moisture.png/.pdf
        q4_fig3_moisture_curves.png/.pdf

说明：
1. 图1读取 file/附件2.xlsx，并与 q4_solver.py 一致采用 PCHIP 构造 R(t)。
2. 图2、图3直接读取 output/q4/result4.xlsx，不重复求解 PDE。
3. result4 中超出当前动态半径 R(t) 的固定空间位置为空值；图2将这些区域保持为空白，
   并用线性插值仅在药材内部重建更细的绘图网格。
4. 当前论文计算得到连续烘干结束时刻为 51.094755 h；可通过 --dry-time-h 修改。
5. 所有图同时保存 600 dpi PNG 与矢量 PDF，适合直接插入 Word 论文。
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Tuple

import numpy as np
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.ticker import MaxNLocator
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator


# ============================================================
# 0. 路径与默认配置
# ============================================================

SCRIPT_DIR = Path(__file__).resolve().parent

# 本脚本预期放在：26MathModel/plot/q4/q4_plot.py
# SCRIPT_DIR.parents[0] = .../26MathModel/plot
# SCRIPT_DIR.parents[1] = .../26MathModel
PROJECT_ROOT = SCRIPT_DIR.parents[1]

DEFAULT_ATTACHMENT2 = "file/附件2.xlsx"
DEFAULT_RESULT4 = "output/q4/result4.xlsx"
DEFAULT_OUTPUT_DIR = "plot/output/q4"

# q4_solver.py 当前连续事件定位结果：183941.1165 s = 51.094755 h
DEFAULT_DRY_TIME_H = 51.094755
TARGET_MOISTURE = 0.15
INITIAL_MOISTURE = 2.55


# ============================================================
# 1. 路径与绘图风格
# ============================================================

def resolve_input_path(path_str: str) -> Path:
    """依次尝试当前工作目录、项目根目录、脚本目录。"""
    p = Path(path_str).expanduser()
    if p.is_absolute():
        return p

    candidates = [
        Path.cwd() / p,
        PROJECT_ROOT / p,
        SCRIPT_DIR / p,
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate.resolve()

    return (PROJECT_ROOT / p).resolve()


def resolve_output_path(path_str: str) -> Path:
    """相对输出路径统一以项目根目录为基准。"""
    p = Path(path_str).expanduser()
    if p.is_absolute():
        return p
    return (PROJECT_ROOT / p).resolve()


def configure_matplotlib() -> None:
    """与项目中 q1/q2 绘图脚本保持一致的中文论文风格。"""
    available_fonts = {f.name for f in font_manager.fontManager.ttflist}
    candidates = [
        "PingFang SC",          # macOS
        "Microsoft YaHei",     # Windows
        "SimHei",              # Windows / Linux
        "Noto Sans CJK SC",    # Linux
        "Source Han Sans SC",  # 思源黑体
        "Arial Unicode MS",
        "DejaVu Sans",
    ]

    for font_name in candidates:
        if font_name in available_fonts:
            plt.rcParams["font.sans-serif"] = [font_name, "DejaVu Sans"]
            break

    plt.rcParams.update(
        {
            "axes.unicode_minus": False,
            "font.size": 10.5,
            "axes.titlesize": 11.5,
            "axes.labelsize": 11,
            "xtick.labelsize": 9.5,
            "ytick.labelsize": 9.5,
            "legend.fontsize": 9.2,
            "figure.dpi": 120,
            "savefig.dpi": 600,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "mathtext.fontset": "stix",
        }
    )


def style_axis(ax: plt.Axes) -> None:
    """统一折线图坐标轴风格。"""
    ax.grid(True, linestyle="--", linewidth=0.55, alpha=0.35)
    ax.tick_params(direction="in", top=True, right=True)
    for spine in ax.spines.values():
        spine.set_linewidth(0.8)


def save_figure(fig: plt.Figure, output_dir: Path, stem: str) -> None:
    """同时保存 600 dpi PNG 与 PDF。"""
    output_dir.mkdir(parents=True, exist_ok=True)
    png_path = output_dir / f"{stem}.png"
    pdf_path = output_dir / f"{stem}.pdf"

    fig.savefig(png_path, dpi=600, bbox_inches="tight", facecolor="white")
    fig.savefig(pdf_path, bbox_inches="tight", facecolor="white")

    print(f"  PNG: {png_path}")
    print(f"  PDF: {pdf_path}")


# ============================================================
# 2. 读取附件2：时间 / 半径
# ============================================================

def read_attachment2(path: Path) -> Tuple[np.ndarray, np.ndarray]:
    """
    读取附件2.xlsx：
        第1列：时间/s
        第2列：药材半径/cm
    """
    if not path.exists():
        raise FileNotFoundError(f"找不到附件2：{path}")

    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        records: list[tuple[float, float]] = []

        for row_index, row in enumerate(
            ws.iter_rows(min_row=2, max_col=2, values_only=True), start=2
        ):
            if row[0] is None and row[1] is None:
                continue
            if row[0] is None or row[1] is None:
                raise ValueError(f"附件2第 {row_index} 行存在缺失数据：{row}")
            records.append((float(row[0]), float(row[1])))
    finally:
        wb.close()

    if len(records) < 2:
        raise ValueError("附件2有效数据不足。")

    arr = np.asarray(records, dtype=float)
    times_s = arr[:, 0]
    radius_cm = arr[:, 1]

    if np.any(np.diff(times_s) <= 0):
        raise ValueError("附件2时间列必须严格递增。")
    if np.any(radius_cm <= 0):
        raise ValueError("附件2半径必须为正数。")

    return times_s, radius_cm


# ============================================================
# 3. 读取 result4.xlsx
# ============================================================

def read_result4(
    path: Path,
) -> Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """
    读取 result4.xlsx。

    返回：
        times_s       : (n_t,)
        distances_cm  : 固定实际距离列，例如 0,0.1,...,1.9
        C_fixed       : (n_t,n_r)，超出当前药材范围的位置为 NaN
        C_surface     : (n_t,)，动态药材表面含水率
    """
    if not path.exists():
        raise FileNotFoundError(f"找不到 result4.xlsx：{path}")

    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)

        try:
            header = next(rows)
        except StopIteration as exc:
            raise ValueError("result4.xlsx 为空。") from exc

        if len(header) < 3:
            raise ValueError("result4.xlsx 列数不足。")

        # 最后一列应为“药材表面”；中间列为固定实际距离。
        if str(header[-1]).strip() != "药材表面":
            raise ValueError(
                "result4.xlsx 最后一列应为“药材表面”；"
                f"当前为 {header[-1]!r}。"
            )

        distances_cm = np.asarray([float(v) for v in header[1:-1]], dtype=float)
        if distances_cm.size == 0:
            raise ValueError("result4.xlsx 未读取到固定实际距离列。")

        times: list[float] = []
        fixed_rows: list[list[float]] = []
        surface_values: list[float] = []

        for row in rows:
            if not row or row[0] is None:
                continue

            time_s = float(row[0])
            fixed_raw = row[1 : 1 + distances_cm.size]
            surface_raw = row[1 + distances_cm.size]

            if surface_raw is None:
                raise ValueError(f"t={time_s} s 的药材表面水分浓度为空。")

            fixed_values = [
                np.nan if v is None else float(v)
                for v in fixed_raw
            ]

            times.append(time_s)
            fixed_rows.append(fixed_values)
            surface_values.append(float(surface_raw))
    finally:
        wb.close()

    if not times:
        raise ValueError("result4.xlsx 没有有效数据行。")

    times_s = np.asarray(times, dtype=float)
    C_fixed = np.asarray(fixed_rows, dtype=float)
    C_surface = np.asarray(surface_values, dtype=float)

    if np.any(np.diff(times_s) <= 0):
        raise ValueError("result4.xlsx 时间列必须严格递增。")
    if C_fixed.shape != (times_s.size, distances_cm.size):
        raise ValueError("result4.xlsx 数据矩阵维度错误。")
    if not np.all(np.isfinite(C_surface)):
        raise ValueError("result4.xlsx 表面水分浓度存在 NaN/Inf。")

    return times_s, distances_cm, C_fixed, C_surface


def get_distance_column(
    distances_cm: np.ndarray,
    C_fixed: np.ndarray,
    target_cm: float,
    atol: float = 1.0e-8,
) -> np.ndarray:
    """取 result4 中指定固定实际距离的列。"""
    indices = np.where(np.isclose(distances_cm, target_cm, atol=atol))[0]
    if indices.size == 0:
        raise ValueError(
            f"result4.xlsx 中找不到 r={target_cm:g} cm 列；"
            f"现有范围 {distances_cm[0]:g}~{distances_cm[-1]:g} cm。"
        )
    return C_fixed[:, int(indices[0])]


# ============================================================
# 4. 构造半径函数 R(t)
# ============================================================

def build_radius_interpolator(
    radius_times_s: np.ndarray,
    radius_cm: np.ndarray,
) -> PchipInterpolator:
    """与 q4_solver.py 保持一致，采用 PCHIP 构造 R(t)。"""
    return PchipInterpolator(radius_times_s, radius_cm, extrapolate=False)


def evaluate_radius_cm(
    times_s: np.ndarray,
    radius_times_s: np.ndarray,
    radius_cm: np.ndarray,
    interpolator: PchipInterpolator,
) -> np.ndarray:
    """
    在附件2实测时间范围内使用 PCHIP；超出后末值保持。
    当前问题4结束时间小于附件2末时刻，因此正常情况下不会用到末值保持。
    """
    result = np.empty_like(times_s, dtype=float)
    first_t = float(radius_times_s[0])
    last_t = float(radius_times_s[-1])

    before = times_s <= first_t
    inside = (times_s > first_t) & (times_s <= last_t)
    after = times_s > last_t

    result[before] = float(radius_cm[0])
    result[inside] = np.asarray(interpolator(times_s[inside]), dtype=float)
    result[after] = float(radius_cm[-1])
    return result


# ============================================================
# 5. 图1：药材半径 R(t) 随时间变化
# ============================================================

def plot_radius_shrinkage(
    radius_times_s: np.ndarray,
    radius_cm: np.ndarray,
    radius_interp: PchipInterpolator,
    dry_time_h: float,
    output_dir: Path,
) -> None:
    dry_time_s = dry_time_h * 3600.0

    # 只展示实际烘干过程，并在结束时刻右侧留少量空白。
    x_end_h = dry_time_h * 1.035
    x_end_s = x_end_h * 3600.0
    dense_s = np.linspace(0.0, min(x_end_s, radius_times_s[-1]), 1200)
    dense_r = evaluate_radius_cm(
        dense_s, radius_times_s, radius_cm, radius_interp
    )

    measured_mask = radius_times_s <= x_end_s + 1.0e-9
    dry_radius_cm = float(
        evaluate_radius_cm(
            np.asarray([dry_time_s]),
            radius_times_s,
            radius_cm,
            radius_interp,
        )[0]
    )

    fig, ax = plt.subplots(figsize=(7.4, 4.75))
    fig.subplots_adjust(left=0.13, right=0.97, bottom=0.16, top=0.90)

    ax.plot(
        dense_s / 3600.0,
        dense_r,
        linewidth=2.0,
        label="PCHIP 插值半径 $R(t)$",
        zorder=2,
    )
    ax.scatter(
        radius_times_s[measured_mask] / 3600.0,
        radius_cm[measured_mask],
        s=18,
        marker="o",
        edgecolors="white",
        linewidths=0.35,
        label="附件2实测点",
        zorder=3,
    )

    ax.axvline(
        dry_time_h,
        linestyle="--",
        linewidth=1.25,
        label=rf"烘干结束 $t_d={dry_time_h:.4f}\,\mathrm{{h}}$",
    )
    ax.scatter(
        [dry_time_h],
        [dry_radius_cm],
        s=42,
        marker="s",
        zorder=4,
    )

    ax.annotate(
        rf"$R(t_d)={dry_radius_cm:.3f}\,\mathrm{{cm}}$",
        xy=(dry_time_h, dry_radius_cm),
        xytext=(-86, 34),
        textcoords="offset points",
        arrowprops=dict(arrowstyle="->", linewidth=0.8),
        fontsize=9.5,
    )

    shrinkage_percent = (1.0 - dry_radius_cm / radius_cm[0]) * 100.0
    ax.text(
        0.02,
        0.05,
        rf"半径由 {radius_cm[0]:.3f} cm 收缩至 {dry_radius_cm:.3f} cm，缩小约 {shrinkage_percent:.1f}%",
        transform=ax.transAxes,
        fontsize=9.2,
        bbox=dict(boxstyle="round,pad=0.25", facecolor="white", alpha=0.82, linewidth=0.6),
    )

    ax.set_title("药材半径随烘干时间的变化")
    ax.set_xlabel(r"时间 $t/\mathrm{h}$")
    ax.set_ylabel(r"药材半径 $R/\mathrm{cm}$")
    ax.set_xlim(0.0, x_end_h)
    ax.set_ylim(min(dense_r.min(), dry_radius_cm) - 0.06, radius_cm[0] + 0.06)
    ax.xaxis.set_major_locator(MaxNLocator(nbins=8))
    style_axis(ax)
    ax.legend(loc="upper right", frameon=True)

    save_figure(fig, output_dir, "q4_fig1_radius_shrinkage")
    plt.close(fig)


# ============================================================
# 6. 图2：移动边界下水分浓度时空云图
# ============================================================

def reconstruct_moving_domain_field(
    times_s: np.ndarray,
    distances_cm: np.ndarray,
    C_fixed: np.ndarray,
    C_surface: np.ndarray,
    radius_at_times_cm: np.ndarray,
    dr_plot_cm: float = 0.01,
) -> Tuple[np.ndarray, np.ndarray]:
    """
    将 result4 的固定实际距离列 + 动态表面值重建到统一细网格。

    对每个时刻：
    - 只在 0 <= r <= R(t) 内做一维线性插值；
    - r > R(t) 始终置为 NaN，不对药材外部做任何外推。
    """
    if dr_plot_cm <= 0:
        raise ValueError("dr_plot_cm 必须大于 0。")

    r_max_cm = max(float(radius_at_times_cm[0]), float(np.nanmax(radius_at_times_cm)))
    n_r = int(np.ceil(r_max_cm / dr_plot_cm)) + 1
    r_plot_cm = np.linspace(0.0, r_max_cm, n_r)
    field = np.full((times_s.size, r_plot_cm.size), np.nan, dtype=float)

    for i in range(times_s.size):
        R = float(radius_at_times_cm[i])
        fixed_values = C_fixed[i]

        # result4 中位于当前药材内部的固定实际位置。
        valid = np.isfinite(fixed_values) & (distances_cm < R - 1.0e-10)
        x = distances_cm[valid]
        y = fixed_values[valid]

        if x.size == 0:
            # 正常 result4 不会出现；保底用中心与表面构造。
            x = np.asarray([0.0], dtype=float)
            y = np.asarray([C_surface[i]], dtype=float)

        # 动态表面是 r=R(t) 的真实边界值。
        x_interp = np.concatenate([x, [R]])
        y_interp = np.concatenate([y, [C_surface[i]]])

        # 防止极少数情况下出现重复横坐标。
        order = np.argsort(x_interp)
        x_interp = x_interp[order]
        y_interp = y_interp[order]
        unique_x, unique_indices = np.unique(x_interp, return_index=True)
        unique_y = y_interp[unique_indices]

        inside = r_plot_cm <= R + 1.0e-12
        field[i, inside] = np.interp(
            r_plot_cm[inside],
            unique_x,
            unique_y,
        )

    return r_plot_cm, field


def plot_moving_boundary_moisture(
    times_s: np.ndarray,
    distances_cm: np.ndarray,
    C_fixed: np.ndarray,
    C_surface: np.ndarray,
    radius_at_times_cm: np.ndarray,
    dry_time_h: float,
    output_dir: Path,
) -> None:
    t_h = times_s / 3600.0
    r_plot_cm, C_plot = reconstruct_moving_domain_field(
        times_s=times_s,
        distances_cm=distances_cm,
        C_fixed=C_fixed,
        C_surface=C_surface,
        radius_at_times_cm=radius_at_times_cm,
        dr_plot_cm=0.01,
    )

    # 让药材外部区域显示为白色。
    cmap = plt.get_cmap("viridis").copy()
    cmap.set_bad("white")

    fig, ax = plt.subplots(figsize=(8.2, 5.0))
    fig.subplots_adjust(left=0.115, right=0.90, bottom=0.15, top=0.90)

    masked = np.ma.masked_invalid(C_plot.T)
    mesh = ax.pcolormesh(
        t_h,
        r_plot_cm,
        masked,
        shading="auto",
        cmap=cmap,
        vmin=float(np.nanmin(C_plot)),
        vmax=float(np.nanmax(C_plot)),
        rasterized=True,
    )

    # C=0.15 干燥阈值前沿。
    if np.nanmin(C_plot) <= TARGET_MOISTURE <= np.nanmax(C_plot):
        contour = ax.contour(
            t_h,
            r_plot_cm,
            C_plot.T,
            levels=[TARGET_MOISTURE],
            colors="white",
            linewidths=1.35,
        )
        ax.clabel(
            contour,
            inline=True,
            fontsize=8.6,
            fmt={TARGET_MOISTURE: r"$C=0.15$"},
        )

    # 移动边界 R(t)。
    ax.plot(
        t_h,
        radius_at_times_cm,
        linewidth=1.8,
        label=r"移动边界 $r=R(t)$",
    )
    ax.axvline(
        dry_time_h,
        linestyle="--",
        linewidth=1.15,
        label=rf"$t_d={dry_time_h:.4f}\,\mathrm{{h}}$",
    )

    cbar = fig.colorbar(mesh, ax=ax, pad=0.025)
    cbar.set_label(r"水分浓度 $C/(\mathrm{kg}\cdot\mathrm{kg}^{-1})$")

    ax.set_title("移动边界下药材内部水分浓度时空分布")
    ax.set_xlabel(r"时间 $t/\mathrm{h}$")
    ax.set_ylabel(r"到药材中心实际距离 $r/\mathrm{cm}$")
    ax.set_xlim(0.0, max(float(t_h[-1]), dry_time_h) * 1.005)
    ax.set_ylim(0.0, float(np.nanmax(radius_at_times_cm)) * 1.01)
    ax.tick_params(direction="in", top=True, right=True)
    for spine in ax.spines.values():
        spine.set_linewidth(0.8)
    ax.legend(loc="upper right", frameon=True)

    ax.text(
        0.97,
        0.06,
        "白色区域：收缩后药材外部",
        transform=ax.transAxes,
        ha="right",
        va="bottom",
        fontsize=9.0,
        bbox=dict(boxstyle="round,pad=0.25", facecolor="white", alpha=0.82, linewidth=0.6),
    )

    save_figure(fig, output_dir, "q4_fig2_moving_boundary_moisture")
    plt.close(fig)


# ============================================================
# 7. 图3：不同径向位置水分浓度随时间变化
# ============================================================

def plot_moisture_curves(
    times_s: np.ndarray,
    distances_cm: np.ndarray,
    C_fixed: np.ndarray,
    C_surface: np.ndarray,
    dry_time_h: float,
    output_dir: Path,
) -> None:
    t_h = times_s / 3600.0

    C0 = get_distance_column(distances_cm, C_fixed, 0.0)
    C05 = get_distance_column(distances_cm, C_fixed, 0.5)
    C10 = get_distance_column(distances_cm, C_fixed, 1.0)

    fig, ax = plt.subplots(figsize=(7.8, 4.9))
    fig.subplots_adjust(left=0.12, right=0.97, bottom=0.16, top=0.90)

    ax.plot(t_h, C0, linewidth=2.0, label=r"中心 $r=0\,\mathrm{cm}$")
    ax.plot(t_h, C05, linewidth=1.7, label=r"$r=0.5\,\mathrm{cm}$")
    ax.plot(t_h, C10, linewidth=1.7, label=r"$r=1.0\,\mathrm{cm}$")
    ax.plot(t_h, C_surface, linewidth=1.7, label="动态药材表面")

    ax.axhline(
        TARGET_MOISTURE,
        linestyle="--",
        linewidth=1.25,
        label=r"烘干阈值 $C=0.15\,\mathrm{kg/kg}$",
    )
    ax.axvline(
        dry_time_h,
        linestyle=":",
        linewidth=1.35,
        label=rf"烘干结束 $t_d={dry_time_h:.4f}\,\mathrm{{h}}$",
    )
    ax.scatter(
        [dry_time_h],
        [TARGET_MOISTURE],
        s=45,
        marker="o",
        zorder=5,
    )

    ax.annotate(
        "中心达到全域控制阈值",
        xy=(dry_time_h, TARGET_MOISTURE),
        xytext=(-128, 58),
        textcoords="offset points",
        arrowprops=dict(arrowstyle="->", linewidth=0.8),
        fontsize=9.4,
    )

    ax.set_title("不同径向位置水分浓度随时间的变化")
    ax.set_xlabel(r"时间 $t/\mathrm{h}$")
    ax.set_ylabel(r"水分浓度 $C/(\mathrm{kg}\cdot\mathrm{kg}^{-1})$")
    ax.set_xlim(0.0, max(float(t_h[-1]), dry_time_h) * 1.01)
    ax.set_ylim(0.0, max(INITIAL_MOISTURE, float(np.nanmax(C0))) * 1.045)
    ax.xaxis.set_major_locator(MaxNLocator(nbins=9))
    style_axis(ax)
    ax.legend(loc="upper right", frameon=True, ncol=1)

    save_figure(fig, output_dir, "q4_fig3_moisture_curves")
    plt.close(fig)


# ============================================================
# 8. 主程序
# ============================================================

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="绘制 2026 数模国赛 A 题问题4的 3 张核心论文图。"
    )
    parser.add_argument(
        "--attachment2",
        default=DEFAULT_ATTACHMENT2,
        help=f"附件2路径，默认：{DEFAULT_ATTACHMENT2}",
    )
    parser.add_argument(
        "--result4",
        default=DEFAULT_RESULT4,
        help=f"result4.xlsx 路径，默认：{DEFAULT_RESULT4}",
    )
    parser.add_argument(
        "--output-dir",
        default=DEFAULT_OUTPUT_DIR,
        help=f"图片输出目录，默认：{DEFAULT_OUTPUT_DIR}",
    )
    parser.add_argument(
        "--dry-time-h",
        type=float,
        default=DEFAULT_DRY_TIME_H,
        help=(
            "连续事件定位得到的烘干结束时间/h；"
            f"默认：{DEFAULT_DRY_TIME_H:.6f}"
        ),
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    configure_matplotlib()

    attachment2_path = resolve_input_path(args.attachment2)
    result4_path = resolve_input_path(args.result4)
    output_dir = resolve_output_path(args.output_dir)

    print("[1/4] 读取附件2并建立 PCHIP 半径函数...")
    radius_times_s, radius_cm = read_attachment2(attachment2_path)
    radius_interp = build_radius_interpolator(radius_times_s, radius_cm)

    print("[2/4] 读取 result4.xlsx...")
    times_s, distances_cm, C_fixed, C_surface = read_result4(result4_path)
    radius_at_times_cm = evaluate_radius_cm(
        times_s,
        radius_times_s,
        radius_cm,
        radius_interp,
    )

    dry_time_s = args.dry_time_h * 3600.0
    if dry_time_s > radius_times_s[-1] + 1.0e-9:
        print(
            "警告：烘干结束时刻超过附件2最后实测时刻；"
            "图中后段半径将采用末值保持。"
        )

    print("[3/4] 绘制三张论文图...")
    print("图1：药材半径收缩曲线")
    plot_radius_shrinkage(
        radius_times_s=radius_times_s,
        radius_cm=radius_cm,
        radius_interp=radius_interp,
        dry_time_h=args.dry_time_h,
        output_dir=output_dir,
    )

    print("图2：移动边界下水分浓度时空云图")
    plot_moving_boundary_moisture(
        times_s=times_s,
        distances_cm=distances_cm,
        C_fixed=C_fixed,
        C_surface=C_surface,
        radius_at_times_cm=radius_at_times_cm,
        dry_time_h=args.dry_time_h,
        output_dir=output_dir,
    )

    print("图3：不同径向位置水分浓度曲线")
    plot_moisture_curves(
        times_s=times_s,
        distances_cm=distances_cm,
        C_fixed=C_fixed,
        C_surface=C_surface,
        dry_time_h=args.dry_time_h,
        output_dir=output_dir,
    )

    dry_radius = float(
        evaluate_radius_cm(
            np.asarray([dry_time_s]),
            radius_times_s,
            radius_cm,
            radius_interp,
        )[0]
    )

    print("[4/4] 完成。")
    print(f"连续烘干结束时间：{args.dry_time_h:.6f} h")
    print(f"结束时药材半径：{dry_radius:.4f} cm")
    print(f"输出目录：{output_dir}")


if __name__ == "__main__":
    main()
